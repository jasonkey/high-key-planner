"""Audit generated workbooks against the district's own block calendar.

    python build.py ../students/example.py
    python build.py --blank
    python verify.py

district_block_calendar.json is the 209 events transcribed from the district's
published block calendar. It is an independent source, not something this code
generates, which is the only reason checking against it means anything.
"""
import re, os, json, datetime as dt, sys
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from openpyxl import load_workbook
import planner as p
ALL=[(dt.date.fromisoformat(d),code,s,e)
     for d,code,s,e in json.load(open(os.path.join(HERE,"district_block_calendar.json")))]
PER={"8:20am-9:30am":1,"9:37am-10:47am":2,"12:38pm-1:48pm":4,"1:55pm-3:05pm":5}
ROW={1:7,2:8,4:11,5:12}

def load(fn):
    cells={}; wb=load_workbook(fn)
    for ws in wb.worksheets:
        if ws.title=="KEY": continue
        for col in range(2,7):
            m=re.search(r'(\d+)/(\d+)$',ws.cell(2,col).value)
            mo,dd=int(m.group(1)),int(m.group(2))
            yr=2026 if mo>=9 else 2027
            d=dt.date(yr,mo,dd)
            if d in cells: raise SystemExit("duplicate date %s"%d)
            cells[d]={r:(ws.cell(r,col).value or "") for r in range(3,14)}
    return cells,len(wb.sheetnames)-1

def audit(fn, titles, s2=None, late_days=(), lunch1_days=()):
    cells,nweeks=load(fn); errs=[]
    # coverage
    d=dt.date(2026,9,14)
    while d<=dt.date(2027,6,25):
        if d.weekday()<5 and d not in cells: errs.append(f"missing {d}")
        d+=dt.timedelta(days=1)
    # every district block event, term-aware
    n=0
    for day,sess,a,b in ALL:
        if day not in cells: errs.append(f"{day} off-calendar"); continue
        per=PER[a+"-"+b]; got=cells[day][ROW[per]]; L,num=sess[0],int(sess[1]); n+=1
        want=titles.get(L)
        if isinstance(want,tuple):            # course changes at the semester break
            want = want[0] if day < p.SEM2_START else want[1]
        if want is None:
            if "FREE" not in got and "T BLOCK" not in got: errs.append(f"{day} {sess}: {got!r}")
        elif L=="A" and num in (1,3) and titles.get("A_only"):
            if "FREE" not in got: errs.append(f"{day} {sess}: expected FREE, got {got!r}")
        else:
            if want not in got or sess not in got: errs.append(f"{day} {sess} p{per}: {got!r}")
    # day classification + arrival/dismissal
    for day,c in cells.items():
        num=c[3]
        if   day in p.CLOSED:    ok="NO SCHOOL" in num or "CLOSED" in num
        elif day in p.SNOW_DAYS: ok="SNOW" in num
        elif day==p.LAST_DAY:    ok="LAST DAY" in num
        elif day in p.EXAM_DAYS: ok="EXAM" in num
        elif day in p.MCAS_DAYS: ok="MCAS" in num
        elif day in p.SPECIAL:   ok="SPECIAL" in num
        elif day in p.DAYNUM:    ok=(num=="DAY %d"%p.DAYNUM[day])
        else: ok=False
        if not ok: errs.append(f"{day}: classification {num!r}")
        if day in p.DAYNUM:
            dn=p.DAYNUM[day]
            wa="9:37" if dn in late_days else "8:20"
            if not c[5].startswith(wa): errs.append(f"{day} D{dn}: arrival {c[5]!r}")
            wd="2:15" if dn==6 else "3:05"
            if wd not in c[13]: errs.append(f"{day} D{dn}: dismissal {c[13]!r}")
            if lunch1_days:
                if dn in lunch1_days and "LUNCH 1" not in c[9]: errs.append(f"{day} D{dn}: lunch order")
                if dn not in lunch1_days and "LUNCH 2" not in c[10]: errs.append(f"{day} D{dn}: lunch order")
    # semester-2 course swap
    if s2:
        L,before,after=s2
        d1=dt.date(2027,1,25); d2=dt.date(2027,2,8)
        pre=[v for k,v in cells.items() if k<p.SEM2_START for v in v.values() if before in str(v)]
        post=[v for k,v in cells.items() if k>=p.SEM2_START for v in v.values() if before in str(v)]
        newp=[v for k,v in cells.items() if k>=p.SEM2_START for v in v.values() if after in str(v)]
        if not pre:  errs.append(f"{before} never appears before the semester break")
        if post:     errs.append(f"{before} still appears after Jan 29 ({len(post)}x)")
        if not newp: errs.append(f"{after} never appears after Jan 29")
        print(f"   semester swap: {before} {len(pre)}x before / {len(post)}x after  |  {after} {len(newp)}x after")
    print(f"{fn}\n   {nweeks} weekly tabs · {len(cells)} weekdays · {n} block events checked · ERRORS: {len(errs)}")
    for e in errs[:12]: print("     ",e)
    return len(errs)

tot=0
tot+=audit("High_Key_Planner_Example.xlsx",
      {"A":"SPANISH 2","A_only":True,"B":"ENGLISH 9","C":"STUDIO ART","D":"BIOLOGY",
       "E":"CIVICS","F":"ALGEBRA 1","G":"US HISTORY"},
      late_days=(1,5), lunch1_days=(1,))
cells,nw=load("High_Key_Planner_TEMPLATE.xlsx")
print("High_Key_Planner_TEMPLATE.xlsx\n   %d weekly tabs \u00b7 %d weekdays" % (nw,len(cells)))
print("\nTOTAL ERRORS:",tot)
